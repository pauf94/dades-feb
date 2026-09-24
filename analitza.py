"""
Calcula les mètriques a partir dels JSON que deixa feb_scraper.py i genera:
  - fulls/full_mestre_<competicio>_<temporada>.xlsx  (Partits, Equips, Jugadores, Quintets, Parelles, Trios, Tirs, Paràmetres)
  - resums/<temporada>_<competicio>_J<nn>.json       (el destacat de la jornada, llest per redactar-hi el resum)
 
Ús:
  python analitza.py --comp lf2a --season 2025 --data data-test --jornada 26
  python analitza.py --comp all --season 2026            (agafa l'última jornada disponible)
"""
import argparse, json, math, re, statistics as st
from pathlib import Path
from collections import defaultdict
 
PBP = re.compile(r"^\((.+?)\) (?:(.+?): )?(.+)$")
TEAM_KEYS = ["PT","T2A","T2I","T3A","T3I","TLA","TLI","RO","RD","AS","BR","BP"]
 
def nteam(s): return re.sub(r"\s+", " ", (s or "")).strip().upper()
 
# ---------------------------------------------------------------- càrrega
def load(data_dir, season, comp):
    root = Path(data_dir) / str(season) / comp
    games = []
    for folder in sorted(root.glob("J*")):
        jn = int(folder.name[1:])
        for f in sorted(folder.glob("*.json")):
            if f.name.startswith("_"): continue
            g = json.loads(f.read_text(encoding="utf-8"))
            if "players" not in g: continue
            g["jornada"] = jn
            g["teams"] = [nteam(t) for t in g["teams"]]
            games.append(g)
    return games
 
# ---------------------------------------------------------------- mètriques d'equip
def team_totals(g, ti):
    ps = [p for p in g["players"] if p["team"] == ti]
    t = {k: sum(p.get(k, 0) for p in ps) for k in TEAM_KEYS}
    t["MIN"] = sum(p.get("min", 0) for p in ps)
    return t
 
def possessions(t): return t["T2I"] + t["T3I"] + 0.44 * t["TLI"] + t["BP"] - t["RO"]
 
def team_metrics(t, o):
    tci = t["T2I"] + t["T3I"]; p = possessions(t); op = possessions(o)
    d = dict(PF=t["PT"], PR=o["PT"], poss=p, ritme=(p + op) / 2,
             OER=t["PT"] / p if p else 0, DER=o["PT"] / op if op else 0,
             eFG=(t["T2A"] + 1.5 * t["T3A"]) / tci if tci else 0,
             eFG_riv=(o["T2A"] + 1.5 * o["T3A"]) / (o["T2I"] + o["T3I"]) if (o["T2I"] + o["T3I"]) else 0,
             pRO=t["RO"] / (t["RO"] + o["RD"]) if (t["RO"] + o["RD"]) else 0,
             pRD=t["RD"] / (t["RD"] + o["RO"]) if (t["RD"] + o["RO"]) else 0,
             pBP=t["BP"] / p if p else 0, pBP_riv=o["BP"] / op if op else 0,
             FTr=t["TLI"] / tci if tci else 0, t3r=t["T3I"] / tci if tci else 0,
             pT2=t["T2A"] / t["T2I"] if t["T2I"] else 0, pT3=t["T3A"] / t["T3I"] if t["T3I"] else 0,
             pTL=t["TLA"] / t["TLI"] if t["TLI"] else 0,
             pAS=t["AS"] / (t["T2A"] + t["T3A"]) if (t["T2A"] + t["T3A"]) else 0,
             RO=t["RO"], BP=t["BP"], BR=t["BR"], AS=t["AS"])
    d["net"] = d["OER"] - d["DER"]
    return d
 
def team_rows(games):
    rows = []
    for g in games:
        A, B = team_totals(g, 0), team_totals(g, 1)
        for ti, (t, o) in enumerate([(A, B), (B, A)]):
            m = team_metrics(t, o)
            rows.append(dict(id=g["id"], jornada=g["jornada"], data=g.get("date", ""),
                             equip=g["teams"][ti], rival=g["teams"][1 - ti],
                             local="L" if ti == 0 else "V",
                             resultat="V" if t["PT"] > o["PT"] else "D", **m))
    return rows
 
# ---------------------------------------------------------------- play-by-play
def parse_pbp(g):
    ev = []
    for q, rem, acc in g["pbp"]:
        m = PBP.match(acc)
        if m: ev.append(dict(q=q, rem=rem, team=nteam(m.group(1)), pl=m.group(2), act=m.group(3)))
    return ev
 
def qlen(q): return 600 if q <= 4 else 300
def qstart(q): return (q - 1) * 600 if q <= 4 else 2400 + (q - 5) * 300
def pts_of(act):
    return 2 if act.startswith("TIRO DE 2 ANOTADO") else 3 if act.startswith("TIRO DE 3 ANOTADO") else 1 if act.startswith("TIRO DE 1 ANOTADO") else 0
def short(n):
    p = (n or "").split(" ")
    return (p[0] + p[1]) if len(p) > 1 else (n or "")
 
def walk(g):
    """Reconstrueix quintets, ratxes, temps morts i el marcador segon a segon."""
    ev = parse_pbp(g); teams = g["teams"]
    on = {t: set() for t in teams}
    L = defaultdict(lambda: dict(min=0, pf=0, pa=0, fga=0, fta=0, to=0, orb=0, ofga=0, ofta=0, oto=0, oorb=0))
    sc = {t: 0 for t in teams}; last = 0; miss = None
    runs = []; run = None; tos = []; timeline = []; tpts = {t: 0 for t in teams}
    def key(t): return (t, " / ".join(sorted(short(x) for x in on[t])))
    for e in ev:
        if e["team"] not in sc: continue
        now = qstart(e["q"]) + qlen(e["q"]) - e["rem"]
        if now > last:
            for t in teams:
                if len(on[t]) == 5: L[key(t)]["min"] += (now - last) / 60
            last = now
        act = e["act"]
        if "Entra a pista" in act: on[e["team"]].add(e["pl"]); continue
        if "Sale de pista" in act: on[e["team"]].discard(e["pl"]); continue
        if "Tiempo muerto" in act.lower() or "tiempo muerto" in act.lower():
            tos.append(dict(team=e["team"], now=now, q=e["q"], rem=e["rem"])); continue
        opp = teams[1] if e["team"] == teams[0] else teams[0]
        me = L[key(e["team"])] if len(on[e["team"]]) == 5 else None
        ot = L[key(opp)] if len(on[opp]) == 5 else None
        p = pts_of(act)
        if p:
            sc[e["team"]] += p; tpts[e["team"]] += p
            if me: me["pf"] += p
            if ot: ot["pa"] += p
            timeline.append(dict(now=now, q=e["q"], rem=e["rem"], a=sc[teams[0]], b=sc[teams[1]]))
            if run and run["team"] == e["team"]:
                run["pts"] += p; run["end"] = (e["q"], e["rem"]); run["end_now"] = now
                run["after"] = [sc[teams[0]], sc[teams[1]]]
            else:
                if run and run["pts"] >= 8: runs.append(run)
                run = dict(team=e["team"], opp=opp, pts=p, start=(e["q"], e["rem"]), start_now=now,
                           end=(e["q"], e["rem"]), end_now=now,
                           before=[sc[teams[0]] - (p if e["team"] == teams[0] else 0),
                                   sc[teams[1]] - (p if e["team"] == teams[1] else 0)],
                           after=[sc[teams[0]], sc[teams[1]]], quintet=key(e["team"])[1], quintet_riv=key(opp)[1])
        if act.startswith("TIRO DE 2") or act.startswith("TIRO DE 3"):
            if me: me["fga"] += 1
            if ot: ot["ofga"] += 1
        if act.startswith("TIRO DE 1"):
            if me: me["fta"] += 1
            if ot: ot["ofta"] += 1
        if act.startswith("PÉRDIDA"):
            if me: me["to"] += 1
            if ot: ot["oto"] += 1
        if "FALLADO" in act: miss = e["team"]
        elif "ANOTADO" in act: miss = None
        if act.startswith("REBOTE"):
            if miss == e["team"]:
                if me: me["orb"] += 1
                if ot: ot["oorb"] += 1
            miss = None
    if run and run["pts"] >= 8: runs.append(run)
    for r in runs:
        r["tm"] = [t for t in tos if t["team"] == r["opp"] and r["start_now"] <= t["now"] <= r["end_now"]]
    return dict(lineups=L, runs=runs, timeouts=tos, timeline=timeline, pbp_points=tpts)
 
# ---------------------------------------------------------------- tirs
def zone_of(x, y, three):
    dx, dy = x - 1.575, y - 7.5
    d = math.hypot(dx, dy); a = abs(math.degrees(math.atan2(dy, dx))); side = "E" if dy < 0 else "D"
    if three:
        if x <= 3.2 and a >= 60: return "T_C" + side
        return "T_F" if a <= 22.5 else "T_A" + side
    if d <= 2.0: return "R"
    if x <= 5.8 and abs(dy) <= 2.45: return "Z"
    if a >= 67.5: return "M_B" + side
    return "M_A" + side if a >= 25 else "M_F"
 
PXM = 850 / 28
def to_m(left, top):
    x = left / 100 * 850 / PXM
    y = (top / 100 * 474 - (474 - 15 * PXM) / 2) / PXM
    if x > 14: x, y = 28 - x, 15 - y
    return x, y
 
def shot_zones(g):
    """Assigna zona a cada tir; els N tirs més llunyans de cada jugadora són els seus triples segons l'acta."""
    by = defaultdict(list)
    for s in g["shots"]:
        by[(s["team"], str(int(s["dorsal"])) if str(s["dorsal"]).isdigit() else s["dorsal"])].append(s)
    out = []
    for p in g["players"]:
        k = (p["team"], str(int(p["dorsal"])) if p["dorsal"].isdigit() else p["dorsal"])
        arr = by.get(k, [])
        if not arr: continue
        sh = []
        for s in arr:
            x, y = to_m(s["left"], s["top"])
            d = math.hypot(x - 1.575, y - 7.5)
            metric = max(d / 6.75, (abs(y - 7.5) / 6.6) if x <= 3.2 else 0)
            sh.append(dict(x=x, y=y, made=s["made"], m=metric))
        sh.sort(key=lambda s: -s["m"])
        for i, s in enumerate(sh): s["three"] = i < p.get("T3I", 0)
        for s in sh:
            out.append(dict(equip=g["teams"][p["team"]], jugadora=p["name"], jornada=g["jornada"],
                            zona=zone_of(s["x"], s["y"], s["three"]), encert=bool(s["made"])))
    return out
 
# ---------------------------------------------------------------- agregats de jugadores
PCOLS = ["PT","T2A","T2I","T3A","T3I","TLA","TLI","RO","RD","AS","BR","BP","TF","FC","FR","VAL","PM"]
 
def player_rows(games):
    acc = defaultdict(lambda: dict(PJ=0, TIT=0, MIN=0.0, VALS=[], un=0.0, ud=0.0, orn=0.0, ord=0.0, drn=0.0, drd=0.0,
                                   **{k: 0 for k in PCOLS}))
    for g in games:
        T = [team_totals(g, 0), team_totals(g, 1)]
        for p in g["players"]:
            if p.get("min", 0) <= 0: continue
            o = acc[(g["teams"][p["team"]], p["name"])]
            o["PJ"] += 1; o["TIT"] += 1 if p.get("starter") else 0; o["MIN"] += p["min"]
            for k in PCOLS: o[k] += p.get(k, 0)
            o["VALS"].append(p.get("VAL", 0))
            tm, op = T[p["team"]], T[1 - p["team"]]
            used = p.get("T2I", 0) + p.get("T3I", 0) + 0.44 * p.get("TLI", 0) + p.get("BP", 0)
            o["un"] += used * tm["MIN"] / 5
            o["ud"] += p["min"] * (tm["T2I"] + tm["T3I"] + 0.44 * tm["TLI"] + tm["BP"])
            o["orn"] += p.get("RO", 0) * tm["MIN"] / 5; o["ord"] += p["min"] * (tm["RO"] + op["RD"])
            o["drn"] += p.get("RD", 0) * tm["MIN"] / 5; o["drd"] += p["min"] * (tm["RD"] + op["RO"])
    out = []
    for (equip, nom), o in acc.items():
        tci = o["T2I"] + o["T3I"]
        out.append(dict(equip=equip, jugadora=nom, PJ=o["PJ"], TIT=o["TIT"], MIN=round(o["MIN"], 1),
                        **{k: o[k] for k in PCOLS},
                        us=o["un"] / o["ud"] if o["ud"] else 0,
                        pROj=o["orn"] / o["ord"] if o["ord"] else 0,
                        pRDj=o["drn"] / o["drd"] if o["drd"] else 0,
                        eFG=(o["T2A"] + 1.5 * o["T3A"]) / tci if tci else 0,
                        TS=o["PT"] / (2 * (tci + 0.44 * o["TLI"])) if (tci or o["TLI"]) else 0,
                        VALmax=max(o["VALS"]) if o["VALS"] else 0,
                        VALmitjana=st.mean(o["VALS"]) if o["VALS"] else 0,
                        VALde=st.stdev(o["VALS"]) if len(o["VALS"]) > 1 else 0))
    return sorted(out, key=lambda r: (r["equip"], -r["MIN"]))
 
def combo_rows(games, min_min=(20, 100, 100)):
    Q, C2, C3 = defaultdict(lambda: defaultdict(float)), defaultdict(lambda: defaultdict(float)), defaultdict(lambda: defaultdict(float))
    excluded = []
    for g in games:
        if not g.get("checks", {}).get("pbp_ok"): excluded.append(g["id"]); continue
        w = walk(g)
        for (team, names), v in w["lineups"].items():
            if not names: continue
            pl = names.split(" / ")
            for store, k in ((Q, names),):
                for f, val in v.items(): store[(team, k)][f] += val
            from itertools import combinations
            for n, store in ((2, C2), (3, C3)):
                for c in combinations(pl, n):
                    for f, val in v.items(): store[(team, " / ".join(c))][f] += val
    def fmt(store, mm):
        rows = []
        for (team, names), v in store.items():
            if v["min"] < mm: continue
            poss = v["fga"] + 0.44 * v["fta"] + v["to"] - v["orb"]
            oposs = v["ofga"] + 0.44 * v["ofta"] + v["oto"] - v["oorb"]
            rows.append(dict(equip=team, combinacio=names.replace(" / ", " · "), minuts=round(v["min"], 1),
                             PF=int(v["pf"]), PC=int(v["pa"]), dif=int(v["pf"] - v["pa"]),
                             ORtg=100 * v["pf"] / poss if poss else 0, DRtg=100 * v["pa"] / oposs if oposs else 0))
            rows[-1]["net"] = rows[-1]["ORtg"] - rows[-1]["DRtg"]
        return sorted(rows, key=lambda r: (r["equip"], -r["minuts"]))
    return fmt(Q, min_min[0]), fmt(C2, min_min[1]), fmt(C3, min_min[2]), excluded
 
# ---------------------------------------------------------------- full de càlcul
def write_xlsx(path, comp, season, trows, prows, q, c2, c3, shots, excluded):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter as CL
    from openpyxl.worksheet.table import Table, TableStyleInfo
    F = Font(name="Arial", size=10); FB = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    HDR = PatternFill("solid", start_color="1F3864")
    wb = Workbook(); wb.remove(wb.active)
    def sheet(name, rows, headers, widths=None, pcts=(), dec2=(), dec1=()):
        ws = wb.create_sheet(name)
        for c, h in enumerate(headers, 1):
            x = ws.cell(1, c, h[1]); x.font = FB; x.fill = HDR
            x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r, row in enumerate(rows, 2):
            for c, h in enumerate(headers, 1):
                v = row.get(h[0], "")
                cell = ws.cell(r, c, round(v, 4) if isinstance(v, float) else v); cell.font = F
                if h[0] in pcts: cell.number_format = "0.0%"
                elif h[0] in dec2: cell.number_format = "0.00"
                elif h[0] in dec1: cell.number_format = "0.0"
        ws.freeze_panes = "C2"; ws.row_dimensions[1].height = 28
        for c in range(1, len(headers) + 1):
            ws.column_dimensions[CL(c)].width = (widths or {}).get(headers[c - 1][0], 9)
        if rows:
            t = Table(displayName=re.sub(r"\W", "", name), ref=f"A1:{CL(len(headers))}{len(rows)+1}")
            t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True); ws.add_table(t)
        return ws
    H = [("jornada","J"),("data","Data"),("equip","Equip"),("rival","Rival"),("local","L/V"),("resultat","Res."),
         ("PF","PF"),("PR","PR"),("poss","Poss"),("ritme","Ritme"),("OER","OER"),("DER","DER"),("net","Net"),
         ("eFG","%eFG"),("eFG_riv","%eFG riv."),("pRO","%RO"),("pRD","%RD"),("pBP","%BP"),("pBP_riv","%BP riv."),
         ("FTr","FTr"),("t3r","3PAr"),("pT2","%T2"),("pT3","%T3"),("pTL","%TL"),("pAS","%AS"),
         ("RO","RO"),("BP","BP"),("BR","BR"),("AS","AS"),("id","ID")]
    pct = {"eFG","eFG_riv","pRO","pRD","pBP","pBP_riv","pT2","pT3","pTL","pAS"}
    sheet("Partits", trows, H, {"equip":22,"rival":22,"data":11,"id":10}, pct, {"OER","DER","net","FTr","t3r"}, {"poss","ritme"})
    # Equips: mitjanes i desviacions
    eq = defaultdict(list)
    for r in trows: eq[r["equip"]].append(r)
    erows = []
    for team, rs in eq.items():
        d = dict(equip=team, PJ=len(rs), V=sum(1 for r in rs if r["resultat"] == "V"))
        d["D"] = d["PJ"] - d["V"]; d["pV"] = d["V"] / d["PJ"]
        for k in ["PF","PR","poss","ritme","OER","DER","net","eFG","eFG_riv","pRO","pRD","pBP","pBP_riv","FTr","t3r","pT2","pT3","pTL","pAS","RO","BP","BR","AS"]:
            d[k] = st.mean(r[k] for r in rs)
        for k in ["OER","DER","eFG","pRO","pBP","ritme"]:
            d["DE_" + k] = st.stdev([r[k] for r in rs]) if len(rs) > 1 else 0
        erows.append(d)
    erows.sort(key=lambda d: (-d["V"], d["equip"]))
    EH = [("equip","Equip"),("PJ","PJ"),("V","V"),("D","D"),("pV","%V")] + \
         [(k, n) for k, n in [("PF","PF/P"),("PR","PR/P"),("poss","Poss"),("ritme","Ritme"),("OER","OER"),("DER","DER"),("net","Net"),
                              ("eFG","%eFG"),("eFG_riv","%eFG riv."),("pRO","%RO"),("pRD","%RD"),("pBP","%BP"),("pBP_riv","%BP riv."),
                              ("FTr","FTr"),("t3r","3PAr"),("pT2","%T2"),("pT3","%T3"),("pTL","%TL"),("pAS","%AS"),("RO","RO/P"),("BP","BP/P"),("BR","BR/P"),("AS","AS/P"),
                              ("DE_OER","DE OER"),("DE_DER","DE DER"),("DE_eFG","DE %eFG"),("DE_pRO","DE %RO"),("DE_pBP","DE %BP"),("DE_ritme","DE ritme")]]
    sheet("Equips", erows, EH, {"equip":24}, pct | {"pV","DE_eFG","DE_pRO","DE_pBP"}, {"OER","DER","net","FTr","t3r","DE_OER","DE_DER","DE_ritme"}, {"PF","PR","poss","ritme","RO","BP","BR","AS"})
    PH = [("equip","Equip"),("jugadora","Jugadora"),("PJ","PJ"),("TIT","Tit."),("MIN","MIN"),("PT","PTS"),
          ("T2A","T2A"),("T2I","T2I"),("T3A","T3A"),("T3I","T3I"),("TLA","TLA"),("TLI","TLI"),
          ("RO","RO"),("RD","RD"),("AS","AS"),("BR","BR"),("BP","BP"),("TF","TAP"),("FC","FC"),("FR","FR"),
          ("VAL","VAL"),("PM","+/-"),("VALmax","VAL màx"),("VALmitjana","VAL/P"),("VALde","DE VAL"),
          ("eFG","%eFG"),("TS","%TS"),("us","%ús"),("pROj","%RO"),("pRDj","%RD")]
    sheet("Jugadores", prows, PH, {"equip":22,"jugadora":30}, {"eFG","TS","us","pROj","pRDj"}, {"VALmitjana","VALde"}, {"MIN"})
    CH = [("equip","Equip"),("combinacio","Combinació"),("minuts","Minuts"),("PF","PF"),("PC","PC"),("dif","+/-"),
          ("ORtg","ORtg"),("DRtg","DRtg"),("net","Net")]
    sheet("Quintets", q, CH, {"equip":22,"combinacio":60}, (), {"ORtg","DRtg","net"}, {"minuts"})
    sheet("Parelles", c2, CH, {"equip":22,"combinacio":42}, (), {"ORtg","DRtg","net"}, {"minuts"})
    sheet("Trios", c3, CH, {"equip":22,"combinacio":52}, (), {"ORtg","DRtg","net"}, {"minuts"})
    # Tirs per zona
    zt = defaultdict(lambda: [0, 0]); zl = defaultdict(lambda: [0, 0])
    for s in shots:
        zt[(s["equip"], s["jugadora"], s["zona"])][0] += 1 if s["encert"] else 0
        zt[(s["equip"], s["jugadora"], s["zona"])][1] += 1
        zl[s["zona"]][0] += 1 if s["encert"] else 0; zl[s["zona"]][1] += 1
    zrows = [dict(equip=e, jugadora=j, zona=z, encerts=v[0], intents=v[1], pct=v[0] / v[1] if v[1] else 0,
                  pct_lliga=zl[z][0] / zl[z][1] if zl[z][1] else 0) for (e, j, z), v in sorted(zt.items())]
    sheet("Tirs", zrows, [("equip","Equip"),("jugadora","Jugadora"),("zona","Zona"),("encerts","Encerts"),
                          ("intents","Intents"),("pct","% jugadora"),("pct_lliga","% lliga")],
          {"equip":22,"jugadora":30,"zona":14}, {"pct","pct_lliga"})
    ws = wb.create_sheet("Paràmetres", 0)
    info = [("Competició", comp), ("Temporada", f"{season}-{str(season+1)[-2:]}"),
            ("Partits", len(trows) // 2), ("Generat automàticament", "feb_scraper.py + analitza.py"),
            ("Poss", "T2I + T3I + 0,44·TLI + BP − RO"),
            ("OER / DER", "punts anotats o rebuts per possessió · Net = OER − DER"),
            ("%eFG", "(T2A + 1,5·T3A) / (T2I + T3I)"), ("%TS", "PTS / (2·(TCI + 0,44·TLI))"),
            ("%RO / %RD", "rebots capturats sobre els disponibles (fórmula clàssica)"),
            ("%ús", "possessions de l'equip que acaba la jugadora mentre és a pista"),
            ("Quintets/Parelles/Trios", f"mínims de minuts junts: {20}, {100} i {100}. ORtg i DRtg per 100 possessions"),
            ("Zones de tir", "12 zones; els triples s'ajusten perquè quadrin amb l'acta"),
            ("Partits sense play-by-play", ", ".join(str(x) for x in excluded) or "cap")]
    for r, (k, v) in enumerate(info, 1):
        ws.cell(r, 1, k).font = Font(name="Arial", size=10, bold=True)
        ws.cell(r, 2, v).font = Font(name="Arial", size=10)
    ws.column_dimensions["A"].width = 28; ws.column_dimensions["B"].width = 80
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
 
# ---------------------------------------------------------------- resum de la jornada
def fmt_t(q, rem): return f"{'Q'+str(q) if q<=4 else 'P'+str(q-4)} {rem//60:02d}:{rem%60:02d}"
 
def resum(games, jornada, min_partits_de=4):
    prev = [g for g in games if g["jornada"] < jornada]
    cur = [g for g in games if g["jornada"] == jornada]
    trows_prev, trows_cur = team_rows(prev), team_rows(cur)
    hist = defaultdict(list)
    for r in trows_prev: hist[r["equip"]].append(r)
    lliga = {k: st.mean(r[k] for r in trows_prev) for k in ["OER","eFG","pRO","pBP","ritme","FTr"]} if trows_prev else {}
    out = dict(jornada=jornada, partits=[], desviacions=[], sorpreses=[], ratxes=[], individuals=[],
               finals_ajustats=[], quintets=[], classificacio=[], avis=[])
    # classificació acumulada
    W = defaultdict(int); G = defaultdict(int)
    for r in trows_prev + trows_cur:
        G[r["equip"]] += 1; W[r["equip"]] += 1 if r["resultat"] == "V" else 0
    out["classificacio"] = [dict(equip=t, V=W[t], D=G[t] - W[t]) for t in sorted(W, key=lambda t: (-W[t], t))]
    # jugadores: mitjana pròpia
    pstat = defaultdict(list)
    for g in prev:
        for p in g["players"]:
            if p.get("min", 0) > 0: pstat[p["name"]].append(p)
    for g in cur:
        w = walk(g); teams = g["teams"]
        A, B = team_totals(g, 0), team_totals(g, 1)
        mets = [team_metrics(A, B), team_metrics(B, A)]
        if not g.get("checks", {}).get("pbp_ok"):
            out["avis"].append(f"{g['id']}: el play-by-play no quadra amb l'acta; ratxes i quintets no s'han calculat")
        out["partits"].append(dict(id=g["id"], equips=teams, marcador=g["score"], data=g.get("date",""),
                                   parcials=g.get("quarters", []),
                                   metriques=[{k: round(v, 4) for k, v in m.items()} for m in mets]))
        # desviacions respecte a la pròpia temporada
        for ti, m in enumerate(mets):
            h = hist[teams[ti]]
            for k, nom in [("OER","punts per possessió"),("DER","punts rebuts per possessió"),("eFG","%eFG"),
                           ("pRO","% rebot ofensiu"),("pBP","% pèrdues"),("FTr","ràtio de TL"),("ritme","ritme")]:
                if len(h) >= min_partits_de:
                    mu = st.mean(x[k] for x in h); sd = st.stdev([x[k] for x in h]) if len(h) > 1 else 0
                    z = (m[k] - mu) / sd if sd else 0
                    if abs(z) >= 1.5:
                        out["desviacions"].append(dict(equip=teams[ti], rival=teams[1-ti], metrica=nom,
                                                       valor=round(m[k],4), mitjana=round(mu,4), z=round(z,2)))
                elif lliga and k in lliga and abs(m[k] - lliga[k]) > 0:
                    pass
        # sorpreses
        gi, li = (0, 1) if g["score"][0] > g["score"][1] else (1, 0)
        if mets[gi]["OER"] < mets[li]["OER"]:
            out["sorpreses"].append(dict(id=g["id"], guanyador=teams[gi], tipus="guanya amb menys punts per possessió",
                                         OER=[round(mets[0]["OER"],3), round(mets[1]["OER"],3)]))
        elif mets[gi]["eFG"] < mets[li]["eFG"]:
            out["sorpreses"].append(dict(id=g["id"], guanyador=teams[gi], tipus="guanya tirant pitjor",
                                         eFG=[round(mets[0]["eFG"],3), round(mets[1]["eFG"],3)]))
        # ratxes: 10+ o 8-9 amb temps mort del rival
        for r in w["runs"]:
            if r["pts"] >= 10 or r["tm"]:
                out["ratxes"].append(dict(id=g["id"], equip=r["team"], punts=r["pts"],
                                          de=fmt_t(*r["start"]), a=fmt_t(*r["end"]),
                                          marcador=f"{r['before'][0]}-{r['before'][1]} → {r['after'][0]}-{r['after'][1]}",
                                          quintet=r["quintet"], quintet_rival=r["quintet_riv"],
                                          temps_mort=[fmt_t(t["q"], t["rem"]) for t in r["tm"]]))
        # final ajustat
        tl = w["timeline"]
        if tl:
            last5 = [x for x in tl if x["now"] >= 2100]
            before = [x for x in tl if x["now"] < 2100]
            m5 = (before[-1]["a"] - before[-1]["b"]) if before else 0
            mm = min([abs(x["a"] - x["b"]) for x in last5] + [abs(m5)]) if last5 else abs(m5)
            if abs(m5) <= 5 or mm <= 5:
                out["finals_ajustats"].append(dict(id=g["id"], equips=teams, marcador=g["score"],
                                                   marge_a_5_min=m5, marge_minim=mm))
        # quintets del partit
        for ti, t in enumerate(teams):
            ls = sorted([(n, v) for (tt, n), v in w["lineups"].items() if tt == t and n], key=lambda x: -x[1]["min"])
            if not ls: continue
            elig = [x for x in ls if x[1]["min"] >= 4]
            best = max(elig, key=lambda x: x[1]["pf"] - x[1]["pa"], default=None)
            worst = min(elig, key=lambda x: x[1]["pf"] - x[1]["pa"], default=None)
            f = lambda x: None if not x else dict(quintet=x[0], minuts=round(x[1]["min"],1), dif=int(x[1]["pf"]-x[1]["pa"]))
            out["quintets"].append(dict(id=g["id"], equip=t, mes_usat=f(ls[0]), millor=f(best), pitjor=f(worst)))
        # individuals
        for p in g["players"]:
            if p.get("min", 0) <= 0: continue
            h = pstat.get(p["name"], [])
            vals = [x.get("VAL", 0) for x in h]
            mu = st.mean(vals) if vals else None
            sd = st.stdev(vals) if len(vals) > 1 else 0
            z = (p["VAL"] - mu) / sd if (mu is not None and sd) else None
            dd = sum(1 for v in (p["PT"], p["RO"] + p["RD"], p["AS"]) if v >= 10) >= 2
            tci = p["T2I"] + p["T3I"]
            ts = p["PT"] / (2 * (tci + 0.44 * p["TLI"])) if (tci or p["TLI"]) else 0
            if p["VAL"] >= 18 or p["PT"] >= 18 or dd or (ts >= 0.65 and tci >= 12) or (z is not None and abs(z) >= 1.8 and p["min"] >= 10 and len(vals) >= 5):
                out["individuals"].append(dict(id=g["id"], equip=g["teams"][p["team"]], jugadora=p["name"],
                                               min=p["min"], PT=p["PT"], REB=p["RO"]+p["RD"], AS=p["AS"], BP=p["BP"],
                                               T2=f'{p["T2A"]}/{p["T2I"]}', T3=f'{p["T3A"]}/{p["T3I"]}', TL=f'{p["TLA"]}/{p["TLI"]}',
                                               VAL=p["VAL"], pm=p["PM"], doble_doble=dd, TS=round(ts,3),
                                               VAL_mitjana=round(mu,2) if mu is not None else None,
                                               z=round(z,2) if z is not None else None, partits_previs=len(vals)))
    out["individuals"].sort(key=lambda d: -d["VAL"])
    out["ratxes"].sort(key=lambda d: -d["punts"])
    return out
 
# ---------------------------------------------------------------- principal
COMPS = ["lf", "challenge", "lf2a", "lf2b"]
 
def run_comp(data_dir, season, comp, jornada, outdir="."):
    games = load(data_dir, season, comp)
    if not games:
        print(f"{comp}: no hi ha dades a {data_dir}/{season}/{comp}"); return None
    jn = max(g["jornada"] for g in games) if jornada in (None, "last") else int(jornada)
    trows = team_rows(games); prows = player_rows(games)
    q, c2, c3, excluded = combo_rows(games)
    shots = [s for g in games for s in shot_zones(g)]
    xlsx = Path(outdir) / "fulls" / f"full_mestre_{comp}_{season}.xlsx"
    write_xlsx(xlsx, comp, season, trows, prows, q, c2, c3, shots, excluded)
    r = resum(games, jn)
    rp = Path(outdir) / "resums" / f"{season}_{comp}_J{jn:02d}.json"
    rp.parent.mkdir(parents=True, exist_ok=True)
    rp.write_text(json.dumps(r, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{comp}: {len(games)} partits · full {xlsx} · resum {rp} "
          f"(J{jn}: {len(r['partits'])} partits, {len(r['desviacions'])} desviacions, "
          f"{len(r['ratxes'])} ratxes, {len(r['individuals'])} actuacions)")
    return r
 
if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--comp", default="all")
    ap.add_argument("--season", type=int, default=2026)
    ap.add_argument("--jornada", default="last")
    ap.add_argument("--data", default="data")
    ap.add_argument("--out", default=".")
    a = ap.parse_args()
    for c in (COMPS if a.comp == "all" else [a.comp]):
        run_comp(a.data, a.season, c, a.jornada, a.out)
 


